import logging
import os
import re
import sys
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

import openpyxl  # type: ignore[import-untyped]
from sqlmodel import Session, select

BACKEND_DIR = Path(__file__).resolve().parents[1]
PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from app.core.db import engine, init_db  # noqa: E402
from app.models import Client, Server  # noqa: E402

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

DEFAULT_EXCEL_PATH = PROJECT_ROOT / "Control_Clientes_Consterna_V7_ENTERPRISE.xlsx"
ACTIVE_CLIENTS_SHEET = "👥 Clientes Activos"
CONFIG_SHEET = "⚙️ Configuración"

OFFICIAL_GAMES: list[str] = [
    "Minecraft",
    "Project Zomboid",
    "Palworld",
    "Rust",
    "Valheim",
    "Terraria",
    "Core Keeper",
]

OFFICIAL_PLANS: dict[str, dict[str, float]] = {
    "Plan 4GB": {"usd": 8.50, "cop": 34000.0, "cost_usd": 4.00},
    "Plan 6GB": {"usd": 12.50, "cop": 50000.0, "cost_usd": 6.00},
    "Plan 8GB": {"usd": 16.00, "cop": 64000.0, "cost_usd": 8.00},
    "Plan 12GB": {"usd": 24.00, "cop": 96000.0, "cost_usd": 12.00},
    "Plan 16GB": {"usd": 32.00, "cop": 128000.0, "cost_usd": 16.00},
}

OFFICIAL_EXTRAS: dict[str, dict[str, float]] = {
    "Ninguno (Solo Host)": {"usd": 0.0, "cop": 0.0},
    "Creación de Modpack / Evento ($5 USD)": {"usd": 5.0, "cop": 20000.0},
    "Sysadmin Dedicado Mensual ($6 USD/mes)": {"usd": 6.0, "cop": 24000.0},
    "IP Dedicada Exclusiva ($4 USD/mes)": {"usd": 4.0, "cop": 16000.0},
}

PLAN_ALIASES: dict[str, str] = {
    "tier 4 gb": "Plan 4GB",
    "plan 4gb": "Plan 4GB",
    "plan 4 gb": "Plan 4GB",
    "4gb": "Plan 4GB",
    "4 gb": "Plan 4GB",
    "tier 6 gb": "Plan 6GB",
    "plan 6gb": "Plan 6GB",
    "plan 6 gb": "Plan 6GB",
    "6gb": "Plan 6GB",
    "6 gb": "Plan 6GB",
    "tier 8 gb": "Plan 8GB",
    "plan 8gb": "Plan 8GB",
    "plan 8 gb": "Plan 8GB",
    "8gb": "Plan 8GB",
    "8 gb": "Plan 8GB",
    "tier 12 gb": "Plan 12GB",
    "plan 12gb": "Plan 12GB",
    "plan 12 gb": "Plan 12GB",
    "12gb": "Plan 12GB",
    "12 gb": "Plan 12GB",
    "16 gb dedicado": "Plan 16GB",
    "tier 16 gb": "Plan 16GB",
    "plan 16gb": "Plan 16GB",
    "plan 16 gb": "Plan 16GB",
    "16gb": "Plan 16GB",
    "16 gb": "Plan 16GB",
}

GAME_ALIASES: dict[str, str] = {
    "minecraft": "Minecraft",
    "minecraft paper": "Minecraft",
    "minecraft fabric": "Minecraft",
    "minecraft forge": "Minecraft",
    "minecraft neoforge": "Minecraft",
    "minecraft purpur": "Minecraft",
    "zomboid": "Project Zomboid",
    "project zomboid": "Project Zomboid",
    "palworld": "Palworld",
    "rust": "Rust",
    "valheim": "Valheim",
    "terraria": "Terraria",
    "core keeper": "Core Keeper",
    "corekeeper": "Core Keeper",
}

EXTRA_ALIASES: dict[str, str] = {
    "sin extra": "Ninguno (Solo Host)",
    "ninguno": "Ninguno (Solo Host)",
    "ninguno (solo host)": "Ninguno (Solo Host)",
    "modpack básico": "Creación de Modpack / Evento ($5 USD)",
    "modpack basico": "Creación de Modpack / Evento ($5 USD)",
    "modpack avanzado": "Creación de Modpack / Evento ($5 USD)",
    "creación de modpack / evento ($5 usd)": "Creación de Modpack / Evento ($5 USD)",
    "sysadmin dedicado mensual ($6 usd/mes)": "Sysadmin Dedicado Mensual ($6 USD/mes)",
    "ip dedicada exclusiva ($4 usd/mes)": "IP Dedicada Exclusiva ($4 USD/mes)",
}


def _clean_str(val: Any) -> str | None:
    if val is None:
        return None
    text = str(val).strip()
    if not text or text.lower() in {"nan", "none", "null", "-"}:
        return None
    return text


def normalize_phone(val: Any) -> str | None:
    raw = _clean_str(val)
    if not raw:
        return None
    if raw.endswith(".0"):
        raw = raw[:-2]
    digits = re.sub(r"[^\d]", "", raw)
    return digits if digits else None


def normalize_game(val: Any) -> str:
    raw = _clean_str(val)
    if not raw:
        return "Minecraft"
    return GAME_ALIASES.get(raw.lower(), raw)


def normalize_plan(val: Any) -> str:
    raw = _clean_str(val)
    if not raw:
        return "Plan 4GB"
    return PLAN_ALIASES.get(raw.lower(), raw)


def normalize_extra(val: Any) -> str:
    raw = _clean_str(val)
    if not raw:
        return "Ninguno (Solo Host)"
    return EXTRA_ALIASES.get(raw.lower(), raw)


def normalize_status(val: Any) -> str:
    raw = _clean_str(val)
    if not raw:
        return "Pendiente"
    lowered = raw.lower()
    if any(k in lowered for k in ("encendido", "activo", "active", "al día", "al dia", "online")):
        return "Activo"
    if any(k in lowered for k in ("apagado", "suspendido", "suspended", "vencido", "offline")):
        return "Suspendido"
    if any(k in lowered for k in ("cancelado", "cancelled", "canceled", "baja")):
        return "Cancelado"
    if any(k in lowered for k in ("pendiente", "pending", "cobrar pronto")):
        return "Pendiente"
    return "Pendiente"


def extract_voice_port(notes: str | None) -> str | None:
    if not notes:
        return None
    match = re.search(
        r"(?:voice[_\s]*port|puerto(?:\s+de)?\s+voz|simple\s+voice\s+chat)\s*[:=]\s*(\d+)",
        notes,
        re.IGNORECASE,
    )
    if match:
        return match.group(1)
    return None


def normalize_datetime(val: Any) -> datetime | None:
    if val is None:
        return None
    if isinstance(val, datetime):
        return val if val.tzinfo is not None else val.replace(tzinfo=UTC)
    if isinstance(val, date):
        return datetime(val.year, val.month, val.day, tzinfo=UTC)
    raw = _clean_str(val)
    if not raw:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            dt = datetime.strptime(raw, fmt)
            return dt.replace(tzinfo=UTC)
        except ValueError:
            continue
    return None


def normalize_date(val: Any) -> date | None:
    dt = normalize_datetime(val)
    return dt.date() if dt is not None else None


def _parse_float(val: Any) -> float | None:
    if val is None:
        return None
    if isinstance(val, int | float):
        return float(val)
    raw = _clean_str(val)
    if not raw or raw.startswith("="):
        return None
    cleaned = raw.replace("$", "").replace(",", "").strip()
    try:
        return float(cleaned)
    except ValueError:
        return None


def resolve_pricing(
    plan_name: str,
    extra_service: str,
    currency: str,
    raw_monthly: Any,
    raw_extra_fee: Any,
    raw_host_cost: Any,
) -> tuple[float, float, float]:
    curr_key = "cop" if currency.upper() == "COP" else "usd"
    plan_info = OFFICIAL_PLANS.get(plan_name)
    extra_info = OFFICIAL_EXTRAS.get(extra_service)

    monthly_val = _parse_float(raw_monthly)
    if (monthly_val is None or monthly_val <= 0) and plan_info:
        monthly_val = plan_info[curr_key]

    extra_val = _parse_float(raw_extra_fee)
    if extra_val is None and extra_info:
        extra_val = extra_info[curr_key]

    host_cost_val = _parse_float(raw_host_cost)
    if (host_cost_val is None or host_cost_val <= 0) and plan_info:
        host_cost_val = plan_info["cost_usd"]

    return (
        float(monthly_val or 0.0),
        float(extra_val or 0.0),
        float(host_cost_val or 0.0),
    )


def update_excel_workbook(excel_path: Path | None = None) -> Path:
    """Update '⚙️ Configuración' and '👥 Clientes Activos' in the Excel workbook to v1.9.0 pricing."""
    target_path = excel_path if excel_path is not None else DEFAULT_EXCEL_PATH
    if not target_path.exists():
        raise FileNotFoundError(f"No se encontró el archivo Excel en: {target_path}")

    wb = openpyxl.load_workbook(target_path, data_only=False)

    if CONFIG_SHEET in wb.sheetnames:
        ws_cfg = wb[CONFIG_SHEET]
        games_list = [*OFFICIAL_GAMES, "Otro"]
        for idx in range(2, 15):
            ws_cfg[f"A{idx}"].value = games_list[idx - 2] if idx - 2 < len(games_list) else None

        plans_list = [*OFFICIAL_PLANS.keys(), "Custom"]
        for idx in range(2, 15):
            ws_cfg[f"B{idx}"].value = plans_list[idx - 2] if idx - 2 < len(plans_list) else None

        extras_list = [*OFFICIAL_EXTRAS.keys(), "Custom"]
        for idx in range(2, 15):
            ws_cfg[f"C{idx}"].value = extras_list[idx - 2] if idx - 2 < len(extras_list) else None

        # Reference table of official plan prices & base host costs in F1:I6
        ws_cfg["F1"].value = "Plan"
        ws_cfg["G1"].value = "Precio USD"
        ws_cfg["H1"].value = "Precio COP"
        ws_cfg["I1"].value = "Costo Base USD"
        for row_offset, (p_name, p_vals) in enumerate(OFFICIAL_PLANS.items(), start=2):
            ws_cfg[f"F{row_offset}"].value = p_name
            ws_cfg[f"G{row_offset}"].value = p_vals["usd"]
            ws_cfg[f"H{row_offset}"].value = int(p_vals["cop"])
            ws_cfg[f"I{row_offset}"].value = p_vals["cost_usd"]

        # Reference table of official add-on prices in K1:M5
        ws_cfg["K1"].value = "Servicio Extra / Add-on"
        ws_cfg["L1"].value = "Precio USD"
        ws_cfg["M1"].value = "Precio COP"
        for row_offset, (e_name, e_vals) in enumerate(OFFICIAL_EXTRAS.items(), start=2):
            ws_cfg[f"K{row_offset}"].value = e_name
            ws_cfg[f"L{row_offset}"].value = e_vals["usd"]
            ws_cfg[f"M{row_offset}"].value = int(e_vals["cop"])

    if ACTIVE_CLIENTS_SHEET in wb.sheetnames:
        ws = wb[ACTIVE_CLIENTS_SHEET]
        for r in range(2, ws.max_row + 1):
            client_name = _clean_str(ws[f"B{r}"].value)
            if client_name:
                ws[f"D{r}"].value = normalize_game(ws[f"D{r}"].value)
                ws[f"E{r}"].value = normalize_extra(ws[f"E{r}"].value)
                ws[f"F{r}"].value = normalize_plan(ws[f"F{r}"].value)

            ws[f"H{r}"].value = (
                f'=IFERROR(IF(OR(ISBLANK(F{r}), F{r}="Custom"), "", '
                f"VLOOKUP(F{r}, '{CONFIG_SHEET}'!$F$2:$I$6, IF(G{r}=\"COP\", 3, 2), FALSE)), \"\")"
            )
            ws[f"I{r}"].value = (
                f'=IFERROR(IF(OR(ISBLANK(E{r}), E{r}="Custom", E{r}=\'{CONFIG_SHEET}\'!$K$2), 0, '
                f"VLOOKUP(E{r}, '{CONFIG_SHEET}'!$K$2:$M$5, IF(G{r}=\"COP\", 3, 2), FALSE)), 0)"
            )
            ws[f"J{r}"].value = (
                f'=IFERROR(IF(OR(ISBLANK(F{r}), F{r}="Custom"), "", '
                f"VLOOKUP(F{r}, '{CONFIG_SHEET}'!$F$2:$I$6, 4, FALSE)), \"\")"
            )

    wb.save(target_path)
    logger.info("Excel actualizado con precios y fórmulas v1.9.0: %s", target_path.name)
    return target_path


def _build_email(external_id: str | None, full_name: str, row_idx: int) -> str:
    if external_id:
        slug = re.sub(r"[^a-z0-9]+", "-", external_id.lower()).strip("-")
        if slug:
            return f"{slug}@consterna.cloud"
    name_slug = re.sub(r"[^a-z0-9]+", ".", full_name.lower()).strip(".")
    if name_slug:
        return f"{name_slug}@consterna.cloud"
    return f"client.{row_idx}@consterna.cloud"


def _sync_session(ws: Any, session: Session) -> dict[str, int]:
    init_db(session)
    clients_synced = 0
    servers_synced = 0

    for row_idx in range(2, ws.max_row + 1):
        full_name = _clean_str(ws[f"B{row_idx}"].value)
        if not full_name:
            continue

        external_id = _clean_str(ws[f"A{row_idx}"].value)
        whatsapp = normalize_phone(ws[f"C{row_idx}"].value)
        game = normalize_game(ws[f"D{row_idx}"].value)
        extra_service = normalize_extra(ws[f"E{row_idx}"].value)
        plan_name = normalize_plan(ws[f"F{row_idx}"].value)
        currency = (_clean_str(ws[f"G{row_idx}"].value) or "USD").upper()

        monthly_price, one_time_fee, host_cost_usd = resolve_pricing(
            plan_name=plan_name,
            extra_service=extra_service,
            currency=currency,
            raw_monthly=ws[f"H{row_idx}"].value,
            raw_extra_fee=ws[f"I{row_idx}"].value,
            raw_host_cost=ws[f"J{row_idx}"].value,
        )

        payment_method = _clean_str(ws[f"L{row_idx}"].value)
        ip_address = _clean_str(ws[f"M{row_idx}"].value)
        start_dt = normalize_date(ws[f"N{row_idx}"].value)
        expiration_dt = normalize_datetime(ws[f"P{row_idx}"].value)
        status_val = normalize_status(ws[f"S{row_idx}"].value)
        notes = _clean_str(ws[f"U{row_idx}"].value)

        email = _build_email(external_id, full_name, row_idx)

        client: Client | None = None
        if external_id:
            client = session.exec(
                select(Client).where(Client.external_id == external_id)
            ).first()
        if not client:
            client = session.exec(select(Client).where(Client.email == email)).first()

        if not client:
            client = Client(
                external_id=external_id,
                full_name=full_name,
                email=email,
                whatsapp_number=whatsapp,
            )
            session.add(client)
            session.commit()
            session.refresh(client)
            logger.info("Creado cliente: %s (%s)", client.full_name, client.external_id)
        else:
            client.external_id = external_id or client.external_id
            client.full_name = full_name
            client.whatsapp_number = whatsapp
            session.add(client)
            session.commit()
            session.refresh(client)
            logger.info(
                "Actualizado cliente: %s (%s)", client.full_name, client.external_id
            )
        clients_synced += 1

        voice_port = extract_voice_port(notes)

        server: Server | None = None
        if ip_address:
            # 1. Match by owner and exact IP
            server = session.exec(
                select(Server).where(
                    Server.owner_id == client.id, Server.ip_address == ip_address
                )
            ).first()
            # 2. Match globally by unique IP
            if not server:
                server = session.exec(
                    select(Server).where(Server.ip_address == ip_address)
                ).first()
            # 3. Match pending unassigned server for the same client and game
            if not server:
                server = session.exec(
                    select(Server).where(
                        Server.owner_id == client.id,
                        Server.ip_address.is_(None),  # type: ignore[union-attr]
                        Server.game == game,
                    )
                ).first()
        else:
            # Row without IP: find an unassigned server for this client, game, and plan
            server = session.exec(
                select(Server).where(
                    Server.owner_id == client.id,
                    Server.ip_address.is_(None),  # type: ignore[union-attr]
                    Server.game == game,
                    Server.plan_name == plan_name,
                )
            ).first()

        if not server:
            server = Server(
                owner_id=client.id,
                ip_address=ip_address,
                plan_name=plan_name,
                game=game,
                extra_service=extra_service,
                currency=currency,
                monthly_price=monthly_price,
                one_time_fee=one_time_fee,
                host_cost_usd=host_cost_usd,
                payment_method=payment_method,
                start_date=start_dt,
                voice_port=voice_port,
                notes=notes,
                status=status_val,
                expiration_date=expiration_dt,
            )
            session.add(server)
            session.commit()
            logger.info(
                "  -> Creado servidor para %s (%s - %s)",
                client.full_name,
                game,
                plan_name,
            )
        else:
            server.owner_id = client.id
            server.ip_address = ip_address
            server.plan_name = plan_name
            server.game = game
            server.extra_service = extra_service
            server.currency = currency
            server.monthly_price = monthly_price
            server.one_time_fee = one_time_fee
            server.host_cost_usd = host_cost_usd
            server.payment_method = payment_method
            server.start_date = start_dt
            server.voice_port = voice_port or server.voice_port
            server.notes = notes
            server.status = status_val
            server.expiration_date = expiration_dt
            session.add(server)
            session.commit()
            logger.info(
                "  -> Actualizado servidor de %s (%s - %s)",
                client.full_name,
                game,
                plan_name,
            )
        servers_synced += 1

    return {"clients_synced": clients_synced, "servers_synced": servers_synced}


def import_excel(
    excel_path: Path | str | None = None, session: Session | None = None
) -> dict[str, int]:
    env_override = os.environ.get("CONSTERNA_EXCEL_PATH")
    if excel_path is not None:
        target_path = Path(excel_path)
    elif env_override:
        target_path = Path(env_override)
    else:
        target_path = DEFAULT_EXCEL_PATH

    if not target_path.exists():
        logger.error("No se encontró el archivo Excel en: %s", target_path)
        return {"clients_synced": 0, "servers_synced": 0}

    logger.info("Iniciando importación desde Excel: %s", target_path.name)
    wb = openpyxl.load_workbook(target_path, data_only=True)
    ws = (
        wb[ACTIVE_CLIENTS_SHEET]
        if ACTIVE_CLIENTS_SHEET in wb.sheetnames
        else wb.active
    )

    if session is not None:
        result = _sync_session(ws, session)
    else:
        with Session(engine) as db_session:
            result = _sync_session(ws, db_session)

    logger.info(
        "Importación completada con éxito (%d clientes, %d servidores).",
        result["clients_synced"],
        result["servers_synced"],
    )
    return result


if __name__ == "__main__":
    cli_path = Path(sys.argv[1]) if len(sys.argv) > 1 else None
    update_excel_workbook(cli_path)
    import_excel(cli_path)
